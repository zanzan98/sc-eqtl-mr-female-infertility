# -*- coding: utf-8 -*-
"""
s63_precheck.py  —— 只读预检（不写任何产物）
1) 核验底本 docx / 现有 v4.md 的存在与 md5
2) 核验 4 处锚点在 docx 中唯一
3) 模拟 s63 的 md 同步，检查 md<->docx 一致性是否会通过
日志：logs/s63_precheck.log
"""
import io
import os
import re
import hashlib

ROOT = r"D:/endometriosis_project/11_sc_eqtl_mr_project"
BASE_DOCX = os.path.join(ROOT, "41_论文初稿_带图_v3.docx")
MD = os.path.join(ROOT, "42_论文初稿_带图_v4.md")
LOG = os.path.join(ROOT, "logs", "s63_precheck.log")

import atexit
# 无论中途是否异常，退出时都把日志落盘（PowerShell 控制台不可靠）
atexit.register(lambda: io.open(LOG, "w", encoding="utf-8", newline="\n").write("\n".join(L) + "\n"))

from docx import Document

JOBS = [
    ("I1", "记忆 B 细胞的复核状态因而有待更大规模的单细胞队列确认。"),
    ("I2", "而 CDC42 更可能是这一区域效应在特定免疫细胞中的“标签”。"),
    ("I3", "两层共同解释了 WNT4 在全血可检出却在单细胞血液面板中缺席这一格局。"),
    ("I4", "共享成分由区域 lead 标记的单倍型承载，而非被检验基因自身的 eQTL。"),
]
MD_P3_HEAD = "最显著的观察之一是 CDC42 的细胞类型特异性"
MD_P4_HEAD = "然而，同一区域内的精细定位把结局可信集置于 WNT4 基因体内"

L = []
def p(*a):
    s = " ".join(str(x) for x in a)
    L.append(s)
    try:
        print(s)
    except UnicodeEncodeError:
        # 控制台为 GBK：退化为纯 ASCII 转义，保证不抛异常
        print(s.encode("ascii", "backslashreplace").decode("ascii"))


p("== s63 预检 ==")
for f in (BASE_DOCX, MD):
    if os.path.exists(f):
        b = open(f, "rb").read()
        p("EXIST %-46s %9d B  md5=%s" % (os.path.basename(f), len(b), hashlib.md5(b).hexdigest()))
    else:
        p("MISSING %s" % f)
        raise SystemExit(1)

doc = Document(BASE_DOCX)
paras = [x for x in doc.paragraphs if x.text.strip()]
p("底本非空段落 = %d" % len(paras))

p("-- 锚点唯一性（docx）--")
ok = True
for tag, tail in JOBS:
    hits = [x for x in doc.paragraphs if x.text.strip().endswith(tail)]
    p("  %s 命中 %d 次" % (tag, len(hits)))
    if len(hits) != 1:
        ok = False
        for h in hits[:3]:
            p("      HIT> %s" % h.text.strip()[-60:])
    else:
        p("      锚点全文: %s" % hits[0].text.strip())

p("-- md 定位锚 --")
md = io.open(MD, encoding="utf-8", newline="").read()
lines = md.split("\n")
for name, head in (("P3", MD_P3_HEAD), ("P4", MD_P4_HEAD)):
    idx = [i for i, s in enumerate(lines) if s.startswith(head)]
    p("  %s 起始行命中 %d 次 idx=%s" % (name, len(idx), idx))
    for i in idx:
        p("      L%d len=%d : %s…" % (i, len(lines[i]), lines[i][:60]))

p("-- 模拟 md 同步后的 md<->docx 一致性 --")
base_texts = [x.text.strip() for x in paras]
# 模拟：docx = base + 4 插入（按锚点顺序）
sim = list(base_texts)
for tag, tail in JOBS:
    a_idx = [i for i, t in enumerate(sim) if t.endswith(tail)]
    assert len(a_idx) == 1, tag
    sim.insert(a_idx[0] + 1, "SIM_%s" % tag)  # 占位即可（文本不参与一致性）
# 用真实插入文本更准确：直接读 s63 的 I1..I4
import importlib.util
spec = importlib.util.spec_from_file_location("s63", os.path.join(ROOT, "scripts", "s63_build_v4_from_docx.py"))
m = importlib.util.module_from_spec(spec)
# 只取常量，不执行 main（s63 的 main 在 __main__ 保护内）
spec.loader.exec_module(m)
real = {"I1": m.I1, "I2": m.I2, "I3": m.I3, "I4": m.I4}
now = list(base_texts)
for tag, tail in JOBS:
    a_idx = [i for i, t in enumerate(now) if t.endswith(tail)]
    assert len(a_idx) == 1, tag
    now.insert(a_idx[0] + 1, real[tag].replace("\n", "").strip())
p("  模拟 docx 非空段落 = %d（底本 %d + 4）" % (len(now), len(base_texts)))

# 9 段切分（base_texts 里 §4 第 3 段起）
star_cands = [i for i, t in enumerate(base_texts) if t.startswith(MD_P3_HEAD)]
p("  base_texts 中 P3 起始索引 = %s" % star_cands)
if len(star_cands) == 1:
    st = star_cands[0]
    new9 = base_texts[st:st + 9]
    p("  docx 9 段首句判定：new9[0] 起 P3head=%s / new9[4] 起 P4head=%s"
      % (new9[0].startswith(MD_P3_HEAD), new9[4].startswith(MD_P4_HEAD)))
    for i, t in enumerate(new9):
        p("     [%d] %s…" % (i, t[:46]))
    # 模拟替换
    sim_lines = list(lines)
    p3_idx = [i for i, s in enumerate(sim_lines) if s.startswith(MD_P3_HEAD)]
    p4_idx = [i for i, s in enumerate(sim_lines) if s.startswith(MD_P4_HEAD)]
    sim_lines[p3_idx[0]] = "\n\n".join(new9[:4])
    sim_lines[p4_idx[0]] = "\n\n".join(new9[4:])
    sim_md = "\n".join(sim_lines)

    def norm(s):
        # 剔除 markdown 强调/代码标记与所有空白
        s = re.sub(r"[*`_>#\[\]]", "", s)
        return re.sub(r"\s+", "", s)

    md_flat = norm(sim_md)
    miss_raw = [t for t in now if re.sub(r"\s+", "", t) not in md_flat]
    miss_norm = [t for t in now if norm(t) not in md_flat]
    p("  模拟 md<->docx 一致性（原始口径）：miss = %d" % len(miss_raw))
    p("  模拟 md<->docx 一致性（剔除 md 标记口径）：miss = %d" % len(miss_norm))
    p("  -- 逐一列出（原始口径 miss，共 %d）--" % len(miss_raw))
    for t in miss_raw:
        okn = norm(t) in md_flat
        p("     [%s] %s" % ("NORM_OK" if okn else "REAL_DIFF", t[:78]))
    p("  -- 剔除标记后仍 miss 的（=真实差异）--")
    for t in miss_norm:
        p("     REAL_DIFF> %s" % t[:90])
    p("  PRECHECK_MD_CONSISTENCY=%s" % ("PASS" if not miss_norm else "FAIL"))

p("PRECHECK_ANCHORS=%s" % ("PASS" if ok else "FAIL"))
io.open(LOG, "w", encoding="utf-8", newline="\n").write("\n".join(L) + "\n")
p("PRECHECK_DONE")

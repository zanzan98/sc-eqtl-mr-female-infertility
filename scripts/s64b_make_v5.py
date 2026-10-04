# -*- coding: utf-8 -*-
"""s64b_make_v5.py —— 由 v4（docx 血统）产出 v5
本轮唯一内容改动 = 裁定 5：删除局限段中的性别分层句。
  OLD: 「；暴露来自外周血免疫细胞而结局为生殖道表型，性别分层与生殖道 eQTL 参考的缺失使这些效应能否迁移至靶组织仍待检验。」
  NEW: 「。」
  （句前分号须改为句号：被删句原与该段第 3 条局限用「；」相连，删句后若保留「；」会成为段中悬空分号）

红线遵守：
  ① 不改任何已闭环结果（tables/ 只读）
  ② ★不得丢失 v4 保住的 9 段手改——逐段断言 99 段中其余 98 段与 v4 逐字相同
  ③ 无无法核实来源的新句（本轮无新增句）
  ④ 不涉及性别分层的任何**新增**描述（本轮为纯删除）

日志：logs/s64b_make_v5.log
"""
import io, os, re, shutil, hashlib
from copy import deepcopy

ROOT = r"D:/endometriosis_project/11_sc_eqtl_mr_project"
BUILD = r"D:/_wb_docx_build"
BASE_DOCX = os.path.join(ROOT, "42_论文初稿_带图_v4.docx")
BASE_MD = os.path.join(ROOT, "42_论文初稿_带图_v4.md")
OUT_DOCX = os.path.join(ROOT, "43_论文初稿_带图_v5.docx")
OUT_DOCX_BUILD = os.path.join(BUILD, "43_论文初稿_带图_v5.docx")
OUT_MD = os.path.join(ROOT, "43_论文初稿_带图_v5.md")
LOG = os.path.join(ROOT, "logs", "s64b_make_v5.log")

from docx import Document
from docx.oxml.ns import qn

OLD_SUB = ("；暴露来自外周血免疫细胞而结局为生殖道表型，性别分层与生殖道 eQTL 参考的缺失"
           "使这些效应能否迁移至靶组织仍待检验。")
NEW_SUB = "。"
MD_START = "；暴露来自外周血免疫细胞而结局为生殖道表型，性别分层与生殖道"
MD_END = "仍待检验。"

HAND9_HEADS = [
    "最显著的观察之一是 CDC42 的细胞类型特异性",
    "在 B 细胞中，CDC42 调控 B 细胞受体信号下游的肌动蛋白重组",
    "在单核细胞一侧，非经典单核细胞沿血管内皮巡逻",
    "这一分布提示，chr1p36.12 在免疫细胞中的信号",
    "然而，同一区域内的精细定位把结局可信集置于 WNT4 基因体内",
    "WNT4 是 Wnt 信号通路的成员，在胚胎期参与苗勒管的形成与分化",
    "在成人期，WNT4 继续参与卵巢功能与子宫内膜周期性重塑",
    "更重要的是，WNT4 的调控异常不仅在着床环节发挥作用",
    "综合来看，chr1p36.12 对不孕症的影响",
]
INSERTS = ["三个信号座在结构上并不对称", "三位点在两层暴露上的判读分层",
           "作为对该归属边界的补充探索", "三个信号座的并列比较为这一分层提供了内部校准"]

NS_WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"

L = []
def p(*a):
    s = " ".join(str(x) for x in a); L.append(s)
    try: print(s)
    except UnicodeEncodeError: print(s.encode("ascii", "backslashreplace").decode("ascii"))


def strip_ws(s):
    return re.sub(r"\s+", "", s)


def run_level_replace(para, old, new):
    """在段落内做 run 级外科替换：只改与删除区间重叠的 run 的 w:t，其余 run 的 rPr 原样保留。"""
    runs = para._p.findall(qn("w:r"))
    texts = ["".join(t.text or "" for t in r.findall(qn("w:t"))) for r in runs]
    full = "".join(texts)
    assert full.count(old) == 1, "目标串在段落内出现 %d 次" % full.count(old)
    i = full.index(old); j = i + len(old)
    new_full = full[:i] + new + full[j:]

    pos = 0
    for r, txt in zip(runs, texts):
        s, e = pos, pos + len(txt)
        pos = e
        if e <= i or s >= j:            # 完全在删除区间之外 → 原样
            continue
        # 计算保留部分；★ NEW 必须插在「跨越删除起点 i 的那个 run」上，
        #   否则会只删不插（曾因此回读不符）
        keep_l = txt[:i - s] if s < i else ""
        if s <= i < e:
            keep_l += new
        keep_r = txt[j - s:] if e > j else ""
        nt = keep_l + keep_r
        ts = r.findall(qn("w:t"))
        for k in ts[1:]:
            r.remove(k)
        if not ts:
            continue
        if nt == "":
            r.remove(ts[0])
        else:
            ts[0].text = nt
            ts[0].set(qn("xml:space"), "preserve")
    return new_full


def main():
    import atexit
    atexit.register(lambda: io.open(LOG, "w", encoding="utf-8", newline="\n").write("\n".join(L) + "\n"))
    base_b = open(BASE_DOCX, "rb").read()
    p("== v5 构建（v4 → v5；唯一内容改动 = 裁定 5 删除性别分层句）==")
    p("底本 %s (%d B / md5 %s)" % (os.path.basename(BASE_DOCX), len(base_b),
                                   hashlib.md5(base_b).hexdigest()))

    doc = Document(BASE_DOCX)
    base_paras = [x for x in doc.paragraphs if x.text.strip()]
    base_texts = [x.text.strip() for x in base_paras]
    p("底本非空段落 = %d" % len(base_texts))
    assert len(base_texts) == 99, "v4 非空段落应为 99，实为 %d" % len(base_texts)

    # 定位唯一含「性别分层」的段落
    hits = [x for x in doc.paragraphs if "性别分层" in x.text]
    assert len(hits) == 1, "含「性别分层」的段落数 = %d（应为 1）" % len(hits)
    para = hits[0]
    # ★ python-docx 每次访问 doc.paragraphs 都新建 Paragraph 包装对象 → 不能用 list.index(对象)
    #   必须按**文本**定位
    cand = [i for i, x in enumerate(base_paras) if x.text.strip() == para.text.strip()]
    assert len(cand) == 1, "按文本定位目标段命中 %d 次" % len(cand)
    idx = cand[0]
    old_txt = para.text
    p("目标段 idx=%d，原长 %d 字" % (idx, len(old_txt)))
    assert "性别分层" in old_txt and "仍待检验" in old_txt

    got = run_level_replace(para, OLD_SUB, NEW_SUB)
    assert got == old_txt.replace(OLD_SUB, NEW_SUB), "run 级替换结果与整串替换不一致"
    assert para.text == got, "回读段落文本与预期不符"
    assert "性别分层" not in para.text, "删除后仍含「性别分层」"
    assert "暴露来自外周血免疫细胞而结局为生殖道表型" not in para.text
    p("删除后该段：")
    p("   …%s" % para.text[-160:])

    doc.save(OUT_DOCX)
    shutil.copy2(OUT_DOCX, OUT_DOCX_BUILD)
    p("docx 落盘：%s (%d B)" % (OUT_DOCX, os.path.getsize(OUT_DOCX)))

    # ---------------- 复核 docx ----------------
    chk = Document(OUT_DOCX)
    now = [x.text.strip() for x in chk.paragraphs if x.text.strip()]
    assert len(now) == 99, "v5 非空段落 %d != 99（本轮为纯文本删除，段数不变）" % len(now)
    diff = [i for i, (a, b) in enumerate(zip(base_texts, now)) if a != b]
    p("与 v4 逐段比对：有差异的段 = %s（应仅 idx=%d）" % (diff, idx))
    assert diff == [idx], "除目标段外还有 %d 段被改动" % (len(diff) - 1)
    p("  其余 98 段逐字相同：True（★9 段手改得以保留）")

    t_all = "\n".join(x.text for x in chk.paragraphs)
    for h in HAND9_HEADS:
        assert h in t_all, "9 段手改缺失：%s" % h[:20]
    p("  9 段手改全部在位：9/9")
    for s in INSERTS:
        assert s in t_all, "v4 插入段缺失：%s" % s[:20]
    p("  v4 的 4 处插入全部在位：4/4")

    ext = []
    for pp in chk.paragraphs:
        for inl in pp._p.findall(".//{%s}inline" % NS_WP):
            e = inl.find("{%s}extent" % NS_WP)
            ext.append((round(int(e.get("cx")) / 360000, 2), round(int(e.get("cy")) / 360000, 2)))
    p("  图片 extent = %s" % ext)
    assert ext == [(16.0, 10.04), (16.0, 14.63), (16.0, 19.51), (16.0, 23.1), (16.0, 21.08)]
    assert chk.sections[0]._sectPr.find(qn("w:lnNumType")) is None, "连续行号未关闭"
    for bad, name in [("$", "美元符"), ("\\", "反斜杠"), ("\ufffd", "替换字符"), ("**", "字面星号")]:
        assert bad not in t_all, "成品含 %s" % name
    p("  图 5 / 行号关闭 / 无越界字符：通过")

    # ---------------- 同步 md ----------------
    md = io.open(BASE_MD, encoding="utf-8", newline="").read()
    assert md.count(MD_START) == 1, "md 中起始锚出现 %d 次" % md.count(MD_START)
    i = md.index(MD_START)
    j = md.index(MD_END, i) + len(MD_END)
    span = md[i:j]
    assert strip_ws(span) == strip_ws(OLD_SUB), "md 目标区间与 docx 目标句不一致：\n%r" % span
    new_md = md[:i] + NEW_SUB + md[j:]
    assert "性别分层" not in new_md, "md 删除后仍含「性别分层」"
    assert strip_ws(new_md).count("性别分层") == 0
    io.open(OUT_MD, "w", encoding="utf-8", newline="\n").write(new_md)
    p("md 落盘：%s (%d B)" % (OUT_MD, os.path.getsize(OUT_MD)))
    p("  md 删除区间 = %r -> %r" % (span[:24] + "…" + span[-8:], NEW_SUB))

    # md ↔ docx 一致性（归一化口径）
    def norm(s):
        s = re.sub(r"[*`_>#\[\]]", "", s)
        return re.sub(r"\s+", "", s)
    md_flat = norm(io.open(OUT_MD, encoding="utf-8").read())
    miss = [t for t in now if norm(t) not in md_flat]
    p("  md↔docx 归一化一致性：未见于 md 的 docx 段落 = %d" % len(miss))
    for t in miss[:3]:
        p("     MISS> %s" % t[:70])
    assert not miss, miss[:1]

    p("")
    p("md5  v4.docx=%s" % hashlib.md5(base_b).hexdigest())
    p("md5  v5.docx=%s" % hashlib.md5(open(OUT_DOCX, "rb").read()).hexdigest())
    p("md5  v5.md  =%s" % hashlib.md5(open(OUT_MD, "rb").read()).hexdigest())
    io.open(LOG, "w", encoding="utf-8", newline="\n").write("\n".join(L) + "\n")
    print("RESULT=PASS")


if __name__ == "__main__":
    import traceback
    try:
        main()
    except BaseException:
        L.append("")
        L.append("==== TRACEBACK ====")
        L.append(traceback.format_exc())
        io.open(LOG, "w", encoding="utf-8", newline="\n").write("\n".join(L) + "\n")
        raise

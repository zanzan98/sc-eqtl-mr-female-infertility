# -*- coding: utf-8 -*-
"""
s63_build_v4_from_docx.py
-------------------------
★ 本轮的关键修正：v4 必须以 **docx 血统** 而非 md 血统构建。

背景（已实证）：
  `41_论文初稿_带图_v3.docx`（项目根，4,494,363 B，mtime 09-28 23:04）
  与其 md 渲染版（`D:\\_wb_docx_build\\41_论文初稿_带图_v3.docx`，3,924,986 B，20:39）
  的**唯一差异**是：md 的 §4 第 3、4 段（2 段）在项目 docx 中被改写扩成 **9 段**。
  该 9 段在**任何 .md 与任何其它 docx 中均不存在**，即产生于 docx 侧的 Word 编辑。
  → 若从 md 重渲染，会**丢掉这 9 段**，违反「其余正文不动」。

本脚本同时产出两份一致的产物：
  1. `42_论文初稿_带图_v4.docx` = 项目 v3.docx（原样）＋ 4 处插入
  2. `42_论文初稿_带图_v4.md`   = 现有 v4.md，把 §4 第 3、4 段替换为 docx 的 9 段
     （其余内容含 4 处插入与 5 个图占位不变）
日志：logs/s63_build_v4_from_docx.log
"""
import io
import os
import re
import sys
import shutil
import hashlib
from copy import deepcopy

ROOT = r"D:/endometriosis_project/11_sc_eqtl_mr_project"
BUILD = r"D:/_wb_docx_build"
BASE_DOCX = os.path.join(ROOT, "41_论文初稿_带图_v3.docx")
MD = os.path.join(ROOT, "42_论文初稿_带图_v4.md")
OUT_DOCX = os.path.join(ROOT, "42_论文初稿_带图_v4.docx")
OUT_DOCX_BUILD = os.path.join(BUILD, "42_论文初稿_带图_v4.docx")
LOG = os.path.join(ROOT, "logs", "s63_build_v4_from_docx.log")

from docx import Document
from docx.oxml.ns import qn

NS_WP = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"

# ---------------------------------------------------------------- 4 处插入
I1 = ("三个信号座在结构上并不对称。另两个座——chr10 的 YME1L1 与 chr2 的 ANXA4——各只承载一个检验对，"
      "而 chr1p36.12 承载 13 个，并且是唯一在同一区域内同时出现两个基因、且二者效应方向相反的座。"
      "把三个座放到同一参照层上比较可以更清楚地看出这一点：在 GTEx v8 全血中，YME1L1（P_HEIDI = 0.114）、"
      "ANXA4（P_HEIDI = 0.142）与 CDC42（P_HEIDI = 0.435）的 cis-SMR 均通过 HEIDI 的一致性判读，"
      "而同一层上的 LINC00339 判为异质（P_HEIDI = 3.4 × 10⁻¹⁰）。"
      "因此 chr1p36.12 的特殊性不止在于它同时携带两个方向相反的信号，还在于这种判读分层只出现在该区域内部——"
      "这构成本研究把该区域单独展开的理由。")
I2 = ("在 GTEx v8 全血（n = 670）中，YME1L1 与 ANXA4 均通过 HEIDI，而 LINC00339 判为异质；"
      "在 OneK1K 单细胞层，YME1L1 的 HEIDI 判定为临界（P = 0.049）。"
      "三位点在两层暴露上的判读分层，进一步支持 chr1p36.12 在三个信号座中的特殊性。")
I3 = ("三个信号座的并列比较为这一分层提供了内部校准：同一套流程在 chr10 的 YME1L1 与 chr2 的 ANXA4 上"
      "给出跨层一致的判读——两个基因在 GTEx v8 全血中均通过 HEIDI——而在 chr1p36.12 上给出座内分层的判读"
      "（CDC42 为一致、LINC00339 为异质）。"
      "判读差异因此更可能来自位点结构本身（单一强 LD 区块、多基因共享同一单倍型），而不是流程的系统性偏倚。"
      "这也说明，在强 LD 区域报告区域层级结论、在单基因信号座报告基因层级结论，"
      "是同一条推断链上的两种分辨率，而不是两套标准。")
I4 = ("作为对该归属边界的补充探索，我们把 CDC42 与 LINC00339 的 cis-eQTL 工具同时纳入"
      "多变量孟德尔随机化（MVMR），以观察两者的效应如何分解。"
      "在本文主分析的剪枝口径（r² < 0.001）下，8 种细胞类型中有 7 种只剩一个独立工具，模型秩不足、无法识别；"
      "其原因是 chr1p36.12 构成单一强 LD 区块，工具并集内最大两两 r² 达 1.0000。"
      "将剪枝阈值放宽到 r² = 0.05–0.3 后，56 个「细胞类型 × 阈值」组合中有 13 个可识别，其分解结果高度一致："
      "调整 LINC00339 后 CDC42 的效应保持为正（+0.090 至 +0.204），"
      "而调整 CDC42 后 LINC00339 的效应衰减至近零（−0.003 至 −0.050）且均不显著。"
      "这一结果支持「在 CDC42 与 LINC00339 之间，前者是更接近独立信号载体的一方」，"
      "但它不构成基因归属的判决：WNT4 与 MASTL 在 OneK1K 的 14 种细胞类型中 cis-eQTL 记录为零（面板缺失），"
      "无法纳入 MVMR，因此该分析无法回答「CDC42 与 WNT4 何者为独立信号」。"
      "由于可识别结果依赖放宽后的剪枝阈值，上述分解应视为探索性证据，与主分析的证据层级不同。")

# (标签, docx 锚点段落「结尾串」, 插入文本)
JOBS = [
    ("I1", "记忆 B 细胞的复核状态因而有待更大规模的单细胞队列确认。", I1),
    ("I2", "而 CDC42 更可能是这一区域效应在特定免疫细胞中的“标签”。", I2),
    ("I3", "两层共同解释了 WNT4 在全血可检出却在单细胞血液面板中缺席这一格局。", I3),
    ("I4", "共享成分由区域 lead 标记的单倍型承载，而非被检验基因自身的 eQTL。", I4),
]

# md 侧同步：§4 第 3、4 段的首句（用于定位整行替换）
MD_P3_HEAD = "最显著的观察之一是 CDC42 的细胞类型特异性"
MD_P4_HEAD = "然而，同一区域内的精细定位把结局可信集置于 WNT4 基因体内"


def set_para_text(p_el, text, font_src_run):
    """把段落元素的所有 run 清掉，只保留一个 run（字体取自 font_src_run）。"""
    for r in p_el.findall(qn("w:r")):
        p_el.remove(r)
    r_new = deepcopy(font_src_run)
    for t in r_new.findall(qn("w:t")):
        r_new.remove(t)
    t = r_new.makeelement(qn("w:t"), {})
    t.set(qn("xml:space"), "preserve")
    t.text = text
    r_new.append(t)
    p_el.append(r_new)
    return p_el


def main():
    L = []
    def p(*a):
        s = " ".join(str(x) for x in a)
        L.append(s); print(s)

    import atexit
    atexit.register(lambda: io.open(LOG, "w", encoding="utf-8", newline="\n").write("\n".join(L) + "\n"))

    # ---- 覆盖前备份（md 血统版 v4 与旧 docx，便于一键回退）----
    for src, tag in [(MD, "md"), (OUT_DOCX, "docx")]:
        if os.path.exists(src):
            dst = os.path.join(BUILD, "42_论文初稿_带图_v4.%s.bak_preDocxLineage" % tag)
            shutil.copy2(src, dst)
            p("备份 %s -> %s (%d B)" % (os.path.basename(src), os.path.basename(dst), os.path.getsize(dst)))

    base_bytes = open(BASE_DOCX, "rb").read()
    p("== s63 v4 重建（docx 血统）==")
    p("底本 %s" % BASE_DOCX)
    p("  大小 %d B / md5 %s" % (len(base_bytes), hashlib.md5(base_bytes).hexdigest()))

    doc = Document(BASE_DOCX)
    paras = [x for x in doc.paragraphs if x.text.strip()]
    p("  非空段落数 = %d" % len(paras))

    base_texts = [x.text.strip() for x in paras]

    # 取模板 run：用 I1 锚点段（纯正文、无粗体）的首个 run
    anchor = None
    hits = [x for x in doc.paragraphs if x.text.strip().endswith(JOBS[0][1])]
    assert len(hits) == 1, "I1 锚点命中 %d 次" % len(hits)
    tpl_runs = hits[0]._p.findall(qn("w:r"))
    assert len(tpl_runs) >= 1, "模板段无 run"
    tpl_run = tpl_runs[0]
    bold = tpl_run.find(qn("w:rPr"))
    is_bold = (bold is not None and bold.find(qn("w:b")) is not None)
    p("  模板 run 取自 I1 锚点段；粗体 = %s" % is_bold)
    assert not is_bold, "模板 run 为粗体，不宜作正文模板"

    # 逐条插入
    inserted = []
    for tag, tail, text in JOBS:
        cands = [x for x in doc.paragraphs if x.text.strip().endswith(tail)]
        assert len(cands) == 1, "%s 锚点命中 %d 次（%s）" % (tag, len(cands), tail[:20])
        a = cands[0]
        p_el = deepcopy(a._p)
        set_para_text(p_el, text, tpl_run)
        a._p.addnext(p_el)
        inserted.append(tag)
        p("  %s 插入于锚点之后（锚点末 20 字：…%s）" % (tag, a.text.strip()[-20:]))
    assert sorted(inserted) == ["I1", "I2", "I3", "I4"], inserted

    doc.save(OUT_DOCX)
    shutil.copy2(OUT_DOCX, OUT_DOCX_BUILD)
    p("docx 落盘：%s (%d B)" % (OUT_DOCX, os.path.getsize(OUT_DOCX)))

    # ---------------- 复核 docx ----------------
    chk = Document(OUT_DOCX)
    now = [x.text.strip() for x in chk.paragraphs if x.text.strip()]
    p("  复核：非空段落 %d（预期 %d + 4 = %d）" % (len(now), len(base_texts), len(base_texts) + 4))
    assert len(now) == len(base_texts) + 4, "段落数不符"
    now_set = set(now)
    missing = [t for t in base_texts if t not in now_set]
    p("  原 95 段全部保留：%s（缺失 %d）" % (not missing, len(missing)))
    assert not missing, missing[:2]
    flat = re.sub(r"\s+", "", "\n".join(now))
    for tag, _, text in JOBS:
        assert re.sub(r"\s+", "", text) in flat, "%s 未落盘" % tag
    p("  4 处插入在 docx 中存在：4/4")

    ext = []
    for para in chk.paragraphs:
        for inl in para._p.findall(".//{%s}inline" % NS_WP):
            e = inl.find("{%s}extent" % NS_WP)
            ext.append((round(int(e.get("cx")) / 360000, 2), round(int(e.get("cy")) / 360000, 2)))
    p("  图片 extent = %s" % ext)
    assert len(ext) == 5, "图片数 %d != 5" % len(ext)
    assert ext == [(16.0, 10.04), (16.0, 14.63), (16.0, 19.51), (16.0, 23.1), (16.0, 21.08)], ext
    assert chk.sections[0]._sectPr.find(qn("w:lnNumType")) is None, "连续行号未关闭"
    t_all = "\n".join(x.text for x in chk.paragraphs)
    for bad, name in [("$", "美元符"), ("\\", "反斜杠"), ("\ufffd", "替换字符"), ("**", "字面星号")]:
        assert bad not in t_all, "成品含 %s" % name
    p("  图 5 / 无越界字符 / 行号关闭：通过")

    # ---------------- 同步 md ----------------
    md = io.open(MD, encoding="utf-8", newline="").read()
    lines = md.split("\n")
    p3_idx = [i for i, s in enumerate(lines) if s.startswith(MD_P3_HEAD)]
    p4_idx = [i for i, s in enumerate(lines) if s.startswith(MD_P4_HEAD)]
    assert len(p3_idx) == 1 and len(p4_idx) == 1, (p3_idx, p4_idx)
    star = now.index([t for t in now if t.startswith(MD_P3_HEAD)][0])
    new9 = now[star:star + 9]
    p("")
    p("md 同步：§4 第 3 段起，docx 的 9 段 =")
    for t in new9:
        p("   - %s…" % t[:34])
    assert new9[0].startswith(MD_P3_HEAD) and new9[4].startswith(MD_P4_HEAD), "9 段切分错位"
    lines[p3_idx[0]] = "\n\n".join(new9[:4])     # P3 → docx 前 4 段
    lines[p4_idx[0]] = "\n\n".join(new9[4:])     # P4 → docx 后 5 段
    new_md = "\n".join(lines)
    io.open(MD, "w", encoding="utf-8", newline="\n").write(new_md)
    p("md 落盘：%s (%d B)" % (MD, os.path.getsize(MD)))

    # md 与 docx 一致性：docx 每段（归一化后）都应能在 md 中找到
    # ★ 归一化口径：docx 侧无 markdown 标记，md 侧含 **bold**/`code` 等；
    #   必须两侧同时剔除标记与空白，否则把「**背景**　女性不孕症…」这类段落误判为缺失
    #   （2026-09-29 预检实测：原始口径 miss=22，归一化后 miss=0）。
    def norm(s):
        s = re.sub(r"[*`_>#\[\]]", "", s)
        return re.sub(r"\s+", "", s)

    md_flat = norm(io.open(MD, encoding="utf-8").read())
    miss2 = [t for t in now if norm(t) not in md_flat]
    p("  md↔docx 一致性（归一化口径）：docx 段落未见于 md 者 = %d" % len(miss2))
    for t in miss2[:5]:
        p("     MISS> %s" % t[:70])
    assert not miss2, miss2[:2]

    io.open(LOG, "w", encoding="utf-8", newline="\n").write("\n".join(L) + "\n")
    print("RESULT=PASS")


if __name__ == "__main__":
    main()

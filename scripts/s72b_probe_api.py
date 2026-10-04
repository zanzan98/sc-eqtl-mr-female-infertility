# -*- coding: utf-8 -*-
"""s72b_probe_api.py —— 探查 python-docx 版本与图片插入 API"""
import io
import docx
import docx.parts.story as st

OUT = r"D:/endometriosis_project/11_sc_eqtl_mr_project/logs/_s72b_probe_api.txt"
L = []
L.append("python-docx version = %s" % docx.__version__)
L.append("docx.parts.story members = %s" % [n for n in dir(st) if not n.startswith("_")])
for cls_name in [n for n in dir(st) if not n.startswith("_")]:
    cls = getattr(st, cls_name)
    if isinstance(cls, type) and hasattr(cls, "get_or_add_image"):
        L.append("  %s.get_or_add_image exists" % cls_name)

from docx.image.image import Image as DImage
F = r"D:/endometriosis_project/11_sc_eqtl_mr_project/figures/"
for n in ["FigS1_instrument_power_gating.png", "FigS2_sensitivity_outcome_concordance.png",
          "FigS3_manhattan_by_celltype.png", "FigS4_discovery_volcano.png"]:
    i = DImage.from_file(F + n)
    L.append("%-44s emu=(%d,%d) px=(%d,%d) dpi=(%s,%s) aspect=%.4f" % (
        n, i.width, i.height, i.px_width, i.px_height, i.horz_dpi, i.vert_dpi,
        i.height / i.width))

# 也探查 Document 上的可用方法
d = docx.Document()
L.append("Document.add_picture = %s" % hasattr(d, "add_picture"))
L.append("doc.part type = %s" % type(d.part))
L.append("doc.part has get_or_add_image = %s" % hasattr(d.part, "get_or_add_image"))

with io.open(OUT, "w", encoding="utf-8", newline="\n") as fh:
    fh.write("\n".join(L) + "\n")
print("\n".join(L))
